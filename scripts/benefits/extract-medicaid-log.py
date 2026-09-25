#!/usr/bin/env python3
"""
COL-773 step 1: turn Jessica's Medicaid Log workbook into normalized JSON for the importer.

  python3 scripts/benefits/extract-medicaid-log.py --in "Medicaid Log.xlsx" --out ~/haven-imports/medicaid-log.json
  python3 scripts/benefits/extract-medicaid-log.py --self-test

The workbook and the JSON contain resident health and financial information: keep both OUTSIDE the
repository (the script refuses an --out path inside it). Facility tabs are read by header name, not
position (tabs differ slightly). Eastside is skipped (sold). Murphy Notes is skipped (replaced by the
owner summary). Nothing is written to any database here.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

FACILITY_TABS = {"Plantation": "The Plantation on Summers", "Homewood": "Homewood Lodge, ALF", "Oakridge": "Oakridge ALF",
                 "Rising Oaks": "Rising Oaks ALF", "Grande Cypress": "Grande Cypress ALF"}
SKIP_TABS = {"Eastside", "Murphy Notes", "Contacts"}
# Header keyword -> board step (first match wins, checked in this order).
HEADER_STEPS = [
    (r"intake.*(email|sent)|notice.*(email|sent)", "intake_emailed"),
    (r"45.?day|intake request", "intake_requested"),
    (r"assessment complete", "assessment_complete"),
    (r"score", "score"),
    (r"3008 requested", "form_3008_requested"),
    (r"3008 returned", "form_3008_returned"),
    (r"app.* request", "app_requested"),
    (r"app.* returned", "app_returned"),
    (r"cares processing", "cares_processing"),
    (r"cares appointment", "cares_appointment"),
    (r"approved|denied", "dcf_decision"),
    (r"waitlist", "waitlist"),
    (r"notes", "notes"),
]
DATE_RE = re.compile(r"(\d{1,2})[/.](\d{1,2})[/.](\d{2,4})")


def to_date(value, as_of: date) -> str | None:
    """A cell becomes the first real date it holds; text without a date becomes None (kept in notes)."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat() if date(2015, 1, 1) <= value <= date(as_of.year + 1, 12, 31) else None
    if isinstance(value, str):
        m = DATE_RE.search(value)
        if m:
            month, day, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
            year = year + 2000 if year < 100 else year
            try:
                return date(year, month, day).isoformat()
            except ValueError:
                return None
    return None


def parse_score(value, as_of: date) -> tuple[int | None, str | None]:
    """'5 / Y' -> 5; '4-Waitlist - Resubmit 10/30' -> (4, next 10/30 on or after as_of)."""
    if value is None:
        return None, None
    text = str(value)
    m = re.match(r"\s*([1-5])\b", text)
    score = int(m.group(1)) if m else None
    reapply = None
    r = re.search(r"resubmit\s+(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", text, re.I)
    if r:
        month, day = int(r.group(1)), int(r.group(2))
        year = int(r.group(3)) if r.group(3) else as_of.year
        year = year + 2000 if year < 100 else year
        try:
            d = date(year, month, day)
            if not r.group(3) and d < as_of:
                d = date(year + 1, month, day)
            reapply = d.isoformat()
        except ValueError:
            reapply = None
    return score, reapply


def split_name(raw: str) -> dict:
    text = re.sub(r"\(.*?\)", "", raw)
    text = re.sub(r"\s+-\s+.*$", "", text).strip()
    if "," in text:
        last, first = [p.strip() for p in text.split(",", 1)]
    else:
        parts = text.split()
        first, last = (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else (text, "")
    return {"log_name": raw.strip(), "first": first, "last": last}


def caseworker_from(notes: str | None) -> str | None:
    if not notes:
        return None
    m = re.search(r"CW:\s*([A-Za-z][A-Za-z .'-]{1,40}?)(?:[.,;\-]|\s{2}|$)", notes)
    return m.group(1).strip() if m else None


def read_facility_tab(ws, facility: str, as_of: date) -> list[dict]:
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=5):
        if row[0].value and str(row[0].value).strip().lower() == "resident name":
            header_row = row[0].row
            break
    if header_row is None:
        return []
    columns: dict[int, str] = {}
    for cell in ws[header_row]:
        name = str(cell.value or "").strip().lower()
        for pattern, key in HEADER_STEPS:
            if name and re.search(pattern, name) and key not in columns.values():
                columns[cell.column - 1] = key
                break
    rows = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        raw_name = values[0] if values else None
        if not isinstance(raw_name, str) or not raw_name.strip() or raw_name.strip().isdigit():
            continue
        entry = {"facility_tab": facility, "facility_name": FACILITY_TABS[facility], **split_name(raw_name), "steps": {}, "extra": []}
        for index, key in columns.items():
            value = values[index] if index < len(values) else None
            if value is None or str(value).strip() == "":
                continue
            if key == "score":
                entry["score"], entry["reapply_on"] = parse_score(value, as_of)
            elif key == "notes":
                entry["notes"] = str(value).strip()[:4000]
            elif key == "waitlist":
                entry["waitlist"] = str(value).strip()[:200]
            else:
                d = to_date(value, as_of)
                if d:
                    entry["steps"][key] = d
                    if key == "dcf_decision" and isinstance(value, str) and re.search(r"denied", value, re.I):
                        entry["decision"] = "Denied"
                    elif key == "dcf_decision" and isinstance(value, str) and re.search(r"approved", value, re.I):
                        entry["decision"] = "Approved"
                else:
                    entry["extra"].append(f"{key}: {str(value).strip()[:200]}")
        entry["caseworker_name"] = caseworker_from(entry.get("notes"))
        if entry["extra"]:
            entry["notes"] = ((entry.get("notes") or "") + " | Log cells: " + "; ".join(entry["extra"]))[:4000].strip(" |")
        del entry["extra"]
        rows.append(entry)
    return rows


def read_completed(ws, as_of: date) -> list[dict]:
    rows, facility = [], None
    for values in ws.iter_rows(values_only=True):
        first = values[0] if values else None
        if isinstance(first, str) and first.strip() in FACILITY_TABS and not any(v not in (None, "") for v in (values[1:4] if values else ())):
            facility = first.strip()
            continue
        if facility is None or not isinstance(first, str) or first.strip().lower() in ("has ltc", ""):
            continue
        plan = values[1] if len(values) > 1 else None
        amount = values[2] if len(values) > 2 else None
        start = to_date(values[3], as_of) if len(values) > 3 else None
        if not isinstance(plan, str) or not plan.strip() or to_date(plan, as_of):
            continue
        cents = int(round(float(amount) * 100)) if isinstance(amount, (int, float)) and amount > 0 else None
        rows.append({"facility_tab": facility, "facility_name": FACILITY_TABS[facility], **split_name(first), "steps": ({"plan_enrolled": start} if start else {}),
                     "plan": plan.strip()[:200], "monthly_cents": cents, "coverage_start": start, "completed": True})
    return rows


def extract(path: str, as_of: date) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    rows, skipped = [], []
    for ws in wb.worksheets:
        if ws.title in SKIP_TABS:
            skipped.append(ws.title)
        elif ws.title in FACILITY_TABS:
            rows.extend(read_facility_tab(ws, ws.title, as_of))
        elif ws.title == "Completed":
            rows.extend(read_completed(ws, as_of))
        else:
            skipped.append(ws.title)
    contacts = []
    if "Contacts" in wb.sheetnames:
        for values in wb["Contacts"].iter_rows(values_only=True):
            cells = [str(v).strip() for v in values if v not in (None, "")]
            if len(cells) >= 2 and re.fullmatch(r"[A-Z][a-z]+ [A-Z][A-Za-z'-]+", cells[0]) and re.search(r"\d{3}[.\-]\d{3}[.\-]\d{4}", cells[1]):
                contacts.append({"name": cells[0], "phone": cells[1]})
    return {"as_of": as_of.isoformat(), "rows": rows, "contacts": contacts, "skipped_tabs": skipped}


def self_test() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Homewood"
    ws.append(["2026-09-15", "HOMEWOOD"])
    ws.append(["Resident Name", "45-Day Notice & Intake Request", "Intake & Notice Email Sent", "Assessment Complete", "Score / Qualify", "3008 Requested",
               "3008 Returned", "Medicaid App Request", "Medicaid App Returned", "Cares Processing Information", "Cares Appointment", "Approved / Denied", "Waitlist", "Notes"])
    ws.append(["Synthetic, Alpha", datetime(2026, 7, 27), datetime(2026, 7, 30), None, "5 / Y", datetime(2026, 8, 12), "08/13/2026 & 09/11/2026", None, None, None, None, "PENDING DCF Review", None, "CW: Tester. POA paperwork"])
    ws.append(["Beta Synthetic - POA", None, None, None, "4-Waitlist - Resubmit 10/30", None, None, None, None, None, None, None, None, None])
    east = wb.create_sheet("Eastside")
    east.append(["Resident Name"])
    east.append(["Should, Skip"])
    done = wb.create_sheet("Completed")
    done.append(["Homewood"])
    done.append(["Has LTC", "Plan", "Amount", "Start Date"])
    done.append(["Gamma Synthetic", "UHC", 1600.0, datetime(2024, 11, 13)])
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "log.xlsx")
        wb.save(path)
        out = extract(path, date(2026, 9, 24))
    rows = {r["log_name"]: r for r in out["rows"]}
    alpha = rows["Synthetic, Alpha"]
    assert alpha["first"] == "Alpha" and alpha["last"] == "Synthetic"
    assert alpha["steps"] == {"intake_requested": "2026-07-27", "intake_emailed": "2026-07-30", "form_3008_requested": "2026-08-12", "form_3008_returned": "2026-08-13"}, alpha["steps"]
    assert alpha["score"] == 5 and alpha["caseworker_name"] == "Tester"
    assert "PENDING DCF Review" in alpha["notes"]
    beta = rows["Beta Synthetic - POA"]
    assert beta["first"] == "Beta" and beta["last"] == "Synthetic" and beta["score"] == 4 and beta["reapply_on"] == "2026-10-30", beta
    gamma = rows["Gamma Synthetic"]
    assert gamma["plan"] == "UHC" and gamma["monthly_cents"] == 160000 and gamma["coverage_start"] == "2024-11-13"
    assert "Should, Skip" not in rows and "Eastside" in out["skipped_tabs"]
    print("self-test PASS")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="source")
    parser.add_argument("--out")
    parser.add_argument("--as-of", default=date.today().isoformat())
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        self_test()
        return
    if not args.source or not args.out:
        parser.error("--in and --out are required")
    repo = os.path.realpath(os.path.join(os.path.dirname(__file__), "..", ".."))
    out_path = os.path.realpath(os.path.expanduser(args.out))
    if out_path.startswith(repo + os.sep):
        sys.exit("Refusing to write resident data inside the repository; choose a path outside it (e.g. ~/haven-imports/).")
    data = extract(args.source, date.fromisoformat(args.as_of))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w") as fh:
        json.dump(data, fh, indent=1)
    by_tab: dict[str, int] = {}
    for row in data["rows"]:
        by_tab[row["facility_tab"] + (" (completed)" if row.get("completed") else "")] = by_tab.get(row["facility_tab"] + (" (completed)" if row.get("completed") else ""), 0) + 1
    print(json.dumps({"rows": len(data["rows"]), "by_tab": by_tab, "contacts": len(data["contacts"]), "skipped_tabs": data["skipped_tabs"]}, indent=1))


if __name__ == "__main__":
    main()
